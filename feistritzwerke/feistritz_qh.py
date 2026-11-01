#!/usr/bin/env python3
"""
feistritz_qh.py - 15-Minuten-Werte aus dem Feistritzwerke-Kundenportal als Excel mit Diagrammen.

Ablauf
  1. Login im Kundenportal (Playwright/Chromium), je Tag "CSV herunterladen" klicken
  2. Alle CSVs einlesen (Semikolon, Dezimalkomma)
  3. Excel schreiben mit den Blättern
       Daten         alle 15-Minuten-Werte (flach, filterbar, Pivot-tauglich)
       Tagessummen   kWh, Spitzenleistung und Intervallanzahl je Tag (Formeln)
       Monatsvergleich / Wochenvergleich
                     je 2 Diagramme; Monate bzw. Wochen per Filter ein-/ausblendbar
       Verlauf       Liniendiagramm über den gesamten Zeitraum
       Tagesvergleich  Tage übereinander gelegt, per Datumsfilter ein-/ausblendbar

Beispiele
  ./feistritz_qh.py --von 07.05.2026 --bis 08.05.2026
  ./feistritz_qh.py --update                 # fehlende Tage seit 01.04.2026 bis zum letzten verfügbaren Tag nachladen
  ./feistritz_qh.py --nur-excel --csv-dir csv   # nur Excel aus vorhandenen CSVs, kein Login
  ./feistritz_qh.py --von 01.05.2026 --bis 31.05.2026 --headed   # Browser sichtbar (Debug)

Zugangsdaten: Umgebungsvariablen FW_USER / FW_PASS, sonst Abfrage im Terminal.
Es wird nichts gespeichert.

Installation
  pip install playwright openpyxl && playwright install chromium
"""

from __future__ import annotations

import argparse
import csv
import getpass
import os
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Paragraph, ParagraphProperties,
                                   RichTextProperties)
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.worksheet.filters import DateGroupItem, FilterColumn, Filters
from openpyxl.worksheet.table import Table, TableStyleInfo

__version__ = "1.5.1"

PORTAL = "https://kundenportal.feistritzwerke.at"
DEFAULT_ANLAGE = "46574"
DEFAULT_ZAEHLPUNKT = "AT0081300000000000000000004657401"
FIRST_DAY = date(2026, 4, 1)   # frühester Tag mit Daten im Portal
MIN_INTERVALS = 92             # ab so vielen Intervallen gilt ein Tag als komplett (DST: 92/96/100)
STOP_AFTER_EMPTY = 3           # --update: nach so vielen leeren Tagen in Folge Ende der Daten
SICHTBAR_TAGE = 14            # Tagesvergleich: so viele Tage sind beim Öffnen eingeblendet
SICHTBAR_MONATE = 3           # Monatsvergleich: so viele Monate sind beim Öffnen eingeblendet
SICHTBAR_WOCHEN = 4           # Wochenvergleich: so viele Wochen sind beim Öffnen eingeblendet
DATE_FMT = "%d.%m.%Y"
TS_FMT = "%d.%m.%Y %H:%M"


# --------------------------------------------------------------------------- Download

def portal_url(anlage: str, zaehlpunkt: str, tag: date) -> str:
    # intervall=4 -> Tag, einheit=2 -> kW (wie in der Portal-URL)
    return (f"{PORTAL}/statistik/anlage/{anlage}"
            f"#zaehlpunkt={zaehlpunkt}+intervall=4+bereich={tag.strftime(DATE_FMT)}+einheit=2")


def daterange(von: date, bis: date):
    d = von
    while d <= bis:
        yield d
        d += timedelta(days=1)


def get_credentials() -> tuple[str, str]:
    user = os.environ.get("FW_USER") or input("Portal-Benutzer: ").strip()
    pw = os.environ.get("FW_PASS") or getpass.getpass("Portal-Passwort: ")
    return user, pw


def login_if_needed(page, user: str, pw: str) -> None:
    """Füllt das Login-Formular aus, falls das Portal eines zeigt."""
    pwd = page.locator("input[type=password]")
    try:
        pwd.first.wait_for(state="visible", timeout=10_000)
    except Exception:
        return  # kein Passwortfeld -> bereits angemeldet
    userfield = page.locator(
        "input[type=email], input[type=text], input:not([type])").first
    userfield.fill(user)
    pwd.first.fill(pw)
    pwd.first.press("Enter")
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1500)  # Redirect nach dem Login abwarten
    if page.locator("input[type=password]").count() and \
            page.locator("input[type=password]").first.is_visible():
        raise SystemExit("Login fehlgeschlagen (Passwortfeld noch sichtbar) - "
                         "Zugangsdaten prüfen oder mit --headed beobachten.")


def download_csvs(tage: list[date], anlage: str, zaehlpunkt: str,
                  outdir: Path, headed: bool, stop_after_empty: int = 0) -> list[Path]:
    from playwright.sync_api import TimeoutError as PWTimeout
    from playwright.sync_api import sync_playwright

    outdir.mkdir(parents=True, exist_ok=True)
    user, pw = get_credentials()
    files: list[Path] = []
    empty = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not headed)
        ctx = browser.new_context(accept_downloads=True, locale="de-AT",
                                  timezone_id="Europe/Vienna")
        page = ctx.new_page()
        page.set_default_timeout(30_000)

        page.goto(portal_url(anlage, zaehlpunkt, tage[0]))
        login_if_needed(page, user, pw)

        for tag in tage:
            target = outdir / f"{zaehlpunkt}_{tag.strftime('%Y%m%d')}_QH.csv"
            # Hash-Wechsel allein lädt die Seite nicht neu (und reload() kollidiert mit
            # Redirects der Seite) -> über about:blank gehen erzwingt einen vollen Ladevorgang
            page.goto("about:blank")
            page.goto(portal_url(anlage, zaehlpunkt, tag), wait_until="load")
            page.wait_for_load_state("networkidle")
            ok = False
            try:
                btn = page.get_by_text("CSV herunterladen")
                btn.wait_for(state="visible", timeout=15_000)
                with page.expect_download(timeout=20_000) as dl:
                    btn.click()
                dl.value.save_as(target)
                ok = bool(read_csv(target))
            except PWTimeout:
                ok = False
            if not ok:
                target.unlink(missing_ok=True)
                empty += 1
                print(f"  {tag.strftime(DATE_FMT)} -> keine Daten verfügbar")
                if stop_after_empty and empty >= stop_after_empty:
                    print(f"  {empty} leere Tage in Folge - Ende der verfügbaren Daten.")
                    break
                continue
            empty = 0
            files.append(target)
            print(f"  {tag.strftime(DATE_FMT)} -> {target.name}")

        browser.close()
    return files


# --------------------------------------------------------------------------- CSV lesen

def read_csv(path: Path) -> list[dict]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    rows = []
    for r in csv.DictReader(text.splitlines(), delimiter=";"):
        if not r.get("Startdatum"):
            continue
        rows.append({
            "start": datetime.strptime(r["Startdatum"], TS_FMT),
            "ende": datetime.strptime(r["Enddatum"], TS_FMT),
            "obis": r.get("OBIS", "").strip(),
            "beschr": r.get("OBIS Kurzbeschreibung", "").strip(),
            "wert": float(r["Wert"].replace(".", "").replace(",", ".")),
            "einheit": r.get("Einheit", "").strip(),
            "qualitaet": r.get("Qualität", "").strip(),
        })
    return rows


def load_all(files: list[Path]) -> list[dict]:
    seen, rows = set(), []
    for f in sorted(files):
        for r in read_csv(f):
            key = (r["obis"], r["start"])
            if key in seen:  # überlappende Exporte nicht doppelt zählen
                continue
            seen.add(key)
            rows.append(r)
    rows.sort(key=lambda r: (r["obis"], r["start"]))
    return rows


def fehlende_tage(csv_dir: Path, von: date, bis: date) -> list[date]:
    """Tage zwischen von und bis, für die keine (vollständigen) Daten in csv_dir liegen."""
    zaehler: dict[date, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for f in csv_dir.glob("*.csv"):
        for r in read_csv(f):
            zaehler[r["start"].date()][r["obis"]] += 1
    return [d for d in daterange(von, bis)
            if not zaehler.get(d) or min(zaehler[d].values()) < MIN_INTERVALS]


# --------------------------------------------------------------------------- Excel

HEAD_FILL = PatternFill("solid", fgColor="8DC63F")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center")


# --------------------------------------------------------------------------- Monats-/Wochenvergleich

GRUEN, GRAU = "8DC63F", "4F6D7A"
PALETTE = ["1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD", "8C564B", "E377C2",
           "7F7F7F", "BCBD22", "17BECF", "393B79", "AD494A"]   # Tagesvergleich: Farben reihum


def _txt(size: int, bold: bool = False) -> RichText:
    cp = CharacterProperties(sz=size, b=bold, solidFill="404040")
    return RichText(bodyPr=RichTextProperties(rot=0, vert="horz", wrap="none"),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _txt_schraeg(size: int) -> RichText:
    """Achsenbeschriftung um 45 Grad gedreht - bricht nie um, auch bei schmalen Säulen."""
    cp = CharacterProperties(sz=size, solidFill="404040")
    return RichText(bodyPr=RichTextProperties(rot=-2700000, vert="horz"),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def style_chart(ch, titel: str, farben: list[str], linie: bool, breite: int = 19050,
                legende: str = "t", plot_breite: float = 0.92) -> None:
    """Einheitlicher, aufgeräumter Chart-Stil: Titel oben, Legende oben, keine Achsentitel
    (Einheit steht im Titel), feste Plot-Fläche -> nichts überlappt."""
    ch.title = titel
    cp = CharacterProperties(sz=1400, b=True, solidFill="404040")
    para = ch.title.tx.rich.p[0]
    para.pPr = ParagraphProperties(defRPr=cp)
    for run in para.r or []:
        run.rPr = cp
    ch.title.overlay = False
    if ch.legend is not None:
        ch.legend.position = legende
        ch.legend.overlay = False
        ch.legend.txPr = _txt(1000)
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.title = None
        ax.txPr = _txt(900)
    ch.x_axis.majorGridlines = None   # nur waagrechte Hilfslinien
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(
        ln=LineProperties(solidFill="D9D9D9", w=6350))
    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge", x=0.06, y=0.18, w=plot_breite, h=0.66))
    for ser, farbe in zip(ch.series, farben):
        if linie:
            ser.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=farbe, w=breite))
        else:
            ser.graphicalProperties = GraphicalProperties(
                solidFill=farbe, ln=LineProperties(solidFill=farbe))


WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
MONTAG_SERIAL = 45292   # 01.01.2024 (Montag) - Basis für die Wochentags-Achse (Format "ddd")


def monat_label(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def woche_label(d: date) -> str:
    jahr, kw, _ = d.isocalendar()
    mo = d - timedelta(days=d.weekday())
    so = mo + timedelta(days=6)
    return f"{jahr}-KW{kw:02d} ({mo:%d.%m.}-{so:%d.%m.})"


def _filter_setzen(ws, spalte: str, kopf: int, ende: int, sichtbar: list[str], alle: list[str],
                   erste: int) -> None:
    """AutoFilter auf einer Spalte; nur 'sichtbar' bleibt eingeblendet (Zeilen + Filterzustand passend)."""
    ws.auto_filter.ref = f"{spalte}{kopf}:{spalte}{ende}"
    ws.auto_filter.filterColumn.append(FilterColumn(colId=0, filters=Filters(filter=list(sichtbar))))
    for k, label in enumerate(alle):
        if label not in sichtbar:
            ws.row_dimensions[erste + k].hidden = True


def add_vergleichsblaetter(wb, rows: list[dict], csv_ist_kwh: bool) -> None:
    """Monats- und Wochenvergleich: je Monat/Woche eine Zeile; der Filter auf der ersten Spalte
    blendet Perioden in BEIDEN Diagrammen ein/aus (Excel zeichnet ausgeblendete Zeilen nicht)."""
    zeiten = [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 15, 30, 45)]
    obis_liste = list(OrderedDict.fromkeys(r["obis"] for r in rows))
    T0, KOPF, ERSTE = 22, 4, 5                      # Spalte V, Kopfzeile, erste Datenzeile

    for idx, obis in enumerate(obis_liste, start=1):
        sub = [r for r in rows if r["obis"] == obis]
        beschr = sub[0]["beschr"] or obis
        suffix = f" {idx}" if len(obis_liste) > 1 else ""

        for art in ("Monat", "Woche"):
            monat = art == "Monat"
            label_fn = monat_label if monat else woche_label
            perioden = sorted({label_fn(r["start"].date()) for r in sub})

            energie: dict = defaultdict(lambda: defaultdict(float))   # periode -> tag/wochentag -> kWh
            ksum: dict = defaultdict(lambda: defaultdict(float))      # periode -> zeitschlüssel -> Summe kW
            kcnt: dict = defaultdict(lambda: defaultdict(int))
            for r in sub:
                kw = r["wert"] * 4 if csv_ist_kwh else r["wert"]
                per = label_fn(r["start"].date())
                z = r["start"].strftime("%H:%M")
                k1 = r["start"].day if monat else r["start"].weekday() + 1
                k2 = z if monat else (r["start"].weekday() + 1, z)
                energie[per][k1] += kw / 4
                ksum[per][k2] += kw
                kcnt[per][k2] += 1

            if monat:
                n1, kopf1 = 31, list(range(1, 32))
                keys2 = zeiten
                x2 = [(int(z[:2]) + int(z[3:]) / 60) / 24 for z in zeiten]       # Tagesbruchteil
                kopf2 = zeiten
            else:
                n1, kopf1 = 7, list(WOCHENTAGE)
                keys2 = [(d, z) for d in range(1, 8) for z in zeiten]
                x2 = [MONTAG_SERIAL + (d - 1) + (int(z[:2]) + int(z[3:]) / 60) / 24 for d, z in keys2]
                kopf2 = [f"{WOCHENTAGE[d - 1]} {z}" for d, z in keys2]
            n2 = len(keys2)
            CL, CS = T0 + 1, T0 + 2                  # Legendentext (mit Summe) und Summe
            C1 = T0 + 3                              # Block 1: Energie je Tag/Wochentag
            C2 = C1 + n1 + 1                         # Block 2: Leistung (eine Leerspalte dazwischen)

            ws = wb.create_sheet(f"{art}svergleich{suffix}" if monat else f"Wochenvergleich{suffix}")
            ws.cell(1, T0, f"{art}e ein-/ausblenden: Pfeil bei „{art}“ (V{KOPF}) anklicken, Haken setzen oder entfernen."
                    ).font = Font(bold=True)
            anzahl = SICHTBAR_MONATE if monat else SICHTBAR_WOCHEN
            ws.cell(2, T0, f"Beim Öffnen sind nur die letzten {anzahl} {art}e sichtbar. "
                           "„Alles auswählen“ im Filter zeigt alle.").font = Font(italic=True, color="7F7F7F")
            ws.cell(2, CS, "Energie (kWh)").font = Font(bold=True)
            ws.cell(2, C2, "Leistung (kW, Ø bei Monat)" if monat else "Leistung (kW)").font = Font(bold=True)
            ws.cell(3, T0, "X-Werte")
            ws.cell(KOPF, T0, art)
            ws.cell(KOPF, CL, "Legende")
            ws.cell(KOPF, CS, "Summe")
            for j, h in enumerate(kopf1):
                ws.cell(KOPF, C1 + j, h)
            for j, (h, x) in enumerate(zip(kopf2, x2)):
                ws.cell(KOPF, C2 + j, h)
                ws.cell(3, C2 + j, x).number_format = "hh:mm" if monat else "[$-407]ddd hh:mm"
            for c in list(range(T0, C1 + n1)) + list(range(C2, C2 + n2)):
                kz = ws.cell(KOPF, c)
                kz.fill, kz.font, kz.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center")
                ws.cell(3, c).font = Font(color="A6A6A6", size=8)
            for k, per in enumerate(perioden):
                rr = ERSTE + k
                ws.cell(rr, T0, per)
                summe = sum(energie[per].values())
                ws.cell(rr, CL, f"{per}  (Σ {summe:.0f} kWh)")
                ws.cell(rr, CS, round(summe, 2)).number_format = "0.00"
                ws.cell(rr, CS).font = Font(bold=True)
                for j, h in enumerate(kopf1):
                    key = h if monat else j + 1
                    if key in energie[per]:
                        ws.cell(rr, C1 + j, round(energie[per][key], 4)).number_format = "0.00"
                for j, key in enumerate(keys2):
                    if kcnt[per][key]:
                        ws.cell(rr, C2 + j, round(ksum[per][key] / kcnt[per][key], 4)).number_format = "0.00"
            ende = ERSTE + len(perioden) - 1
            _filter_setzen(ws, get_column_letter(T0), KOPF, ende, perioden[-anzahl:], perioden, ERSTE)
            ws.column_dimensions[get_column_letter(T0)].width = 24 if not monat else 12
            ws.column_dimensions[get_column_letter(CL)].width = 36
            ws.column_dimensions[get_column_letter(CS)].width = 11
            for c in range(C1, C1 + n1):
                ws.column_dimensions[get_column_letter(c)].width = 7
            for c in range(C2, C2 + n2):
                ws.column_dimensions[get_column_letter(c)].width = 6.5
            ws.sheet_view.zoomScale = 90

            farben = [PALETTE[k % len(PALETTE)] for k in range(len(perioden))]

            # ---- Diagramm 1: Energie je Tag (Linien) bzw. je Wochentag (Säulen)
            c1 = LineChart() if monat else BarChart()
            c1.height, c1.width = 15, 34
            if not monat:
                c1.type, c1.gapWidth = "col", 60
            for k, farbe in enumerate(farben):
                rr = ERSTE + k
                se = Series(Reference(ws, min_col=C1, max_col=C1 + n1 - 1, min_row=rr))
                se.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!${get_column_letter(CL)}${rr}"))
                if monat:
                    se.smooth = False
                    se.marker.symbol, se.marker.size = "circle", 5
                    se.marker.graphicalProperties = GraphicalProperties(
                        solidFill=farbe, ln=LineProperties(solidFill=farbe))
                c1.series.append(se)
            c1.set_categories(Reference(ws, min_col=C1, max_col=C1 + n1 - 1, min_row=KOPF))
            style_chart(c1, f"{beschr} - Energie je {'Tag' if monat else 'Wochentag'} (kWh)", farben,
                        linie=monat, legende="r", plot_breite=0.74)
            ws.add_chart(c1, "A1")

            # ---- Diagramm 2: Leistung über der Tageszeit bzw. über die Woche (Scatter, echte Zeitachse)
            c2 = ScatterChart()
            c2.height, c2.width = 15, 34
            xw = Reference(ws, min_col=C2, max_col=C2 + n2 - 1, min_row=3)
            for k in range(len(perioden)):
                rr = ERSTE + k
                se = Series(Reference(ws, min_col=C2, max_col=C2 + n2 - 1, min_row=rr), xw)
                se.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!${get_column_letter(CL)}${rr}"))
                se.marker.symbol, se.smooth = "none", False
                c2.series.append(se)
            if monat:
                c2.x_axis.scaling.min, c2.x_axis.scaling.max = 0, 1
                c2.x_axis.majorUnit = 2 / 24
                c2.x_axis.number_format = "[hh]:mm"
            else:
                c2.x_axis.scaling.min, c2.x_axis.scaling.max = MONTAG_SERIAL, MONTAG_SERIAL + 6.99
                c2.x_axis.majorUnit = 1
                c2.x_axis.number_format = "[$-407]ddd"
            style_chart(c2, f"{beschr} - " + ("Ø Tagesgang (kW je 15 min)" if monat
                                              else "Leistungsverlauf der Woche (kW)"),
                        farben, linie=True, breite=15875, legende="r", plot_breite=0.74)
            ws.add_chart(c2, "A32")


def build_excel(rows: list[dict], out: Path, csv_ist_kwh: bool) -> None:
    wb = Workbook()

    # ---- Blatt Daten
    ws = wb.active
    ws.title = "Daten"
    head = ["Start", "Ende", "Datum", "Zeit", "OBIS", "Beschreibung",
            "Wert_CSV", "Einheit_CSV", "Leistung_kW", "Energie_kWh", "Qualität",
            "Monat", "Woche", "Tag", "Wochentag"]
    ws.append(head)
    for r in rows:
        kw = r["wert"] * 4 if csv_ist_kwh else r["wert"]
        ws.append([r["start"], r["ende"], r["start"].date(), r["start"].strftime("%H:%M"),
                   r["obis"], r["beschr"], r["wert"], r["einheit"],
                   round(kw, 4), round(kw / 4, 4), r["qualitaet"],
                   monat_label(r["start"].date()), woche_label(r["start"].date()),
                   r["start"].day, r["start"].weekday() + 1])
    n = len(rows) + 1
    style_header(ws, len(head))
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[0].number_format = row[1].number_format = "DD.MM.YYYY HH:MM"
        row[2].number_format = "DD.MM.YYYY"
        row[6].number_format = row[8].number_format = row[9].number_format = "0.000"
    for i, w in enumerate([17, 17, 12, 7, 14, 18, 10, 11, 12, 12, 11, 9, 26, 6, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    tab = Table(displayName="QH", ref=f"A1:{get_column_letter(len(head))}{n}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)

    # Zeilenbereiche je OBIS-Reihe (Zeilen sind nach OBIS, Start sortiert)
    series: "OrderedDict[str, list[int]]" = OrderedDict()
    for i, r in enumerate(rows, start=2):
        series.setdefault(r["obis"], []).append(i)

    # ---- Blatt Tagessummen (Formeln)
    st = wb.create_sheet("Tagessummen")
    st.append(["OBIS", "Datum", "Energie_kWh", "Spitze_kW", "Intervalle", "Achse"])
    style_header(st, 6)
    srow = 2
    summen_zeilen: dict[str, tuple[int, int]] = {}
    for obis in series:
        erste_zeile = srow
        tage = sorted({rows[i - 2]["start"].date() for i in series[obis]})
        for t in tage:
            st.cell(srow, 1, obis)
            st.cell(srow, 2, t).number_format = "DD.MM.YYYY"
            crit = f'Daten!$E:$E,$A{srow},Daten!$C:$C,$B{srow}'
            st.cell(srow, 3, f"=SUMIFS(Daten!$J:$J,{crit})").number_format = "0.00"
            st.cell(srow, 4, f"=_xlfn.MAXIFS(Daten!$I:$I,{crit})").number_format = "0.00"
            st.cell(srow, 5, f"=COUNTIFS({crit})")
            # spärliche X-Achsenbeschriftung für das Tagesenergie-Diagramm (~14 Marken)
            schritt = max(1, -(-len(tage) // 14))
            st.cell(srow, 6, t.strftime("%d.%m.") if (srow - erste_zeile) % schritt == 0 else "")
            srow += 1
        summen_zeilen[obis] = (erste_zeile, srow - 1)
    for i, w in enumerate([16, 12, 13, 11, 11], 1):
        st.column_dimensions[get_column_letter(i)].width = w

    # ---- Diagramme je OBIS-Reihe
    multi = len(series) > 1
    for idx, (obis, zeilen) in enumerate(series.items(), start=1):
        suffix = f" {idx}" if multi else ""
        first, last = zeilen[0], zeilen[-1]
        beschr = rows[first - 2]["beschr"] or obis
        tage = sorted({rows[i - 2]["start"].date() for i in zeilen})

        # Verlauf: Scatter mit echter Zeitachse (Excel-Datumsseriennummer)
        sc = ScatterChart()
        sc.height, sc.width = 12.5, 34
        # Achsen-Raster je nach Zeitspanne: höchstens ~14 Beschriftungen
        spanne = (tage[-1] - tage[0]).days + 1
        if spanne <= 3:
            sc.x_axis.number_format, sc.x_axis.majorUnit = "DD.MM. HH:MM", 0.25   # alle 6 h
        else:
            einheit = next((u for u in (1, 2, 7, 14, 30, 60, 90) if spanne / u <= 14), 90)
            sc.x_axis.number_format, sc.x_axis.majorUnit = "DD.MM.", einheit      # Tage
        sc.legend = None
        sc.x_axis.scaling.min = int(to_excel(rows[first - 2]["start"]))  # ab 00:00 des ersten Tages
        sc.x_axis.scaling.max = int(to_excel(rows[last - 2]["ende"]) + 0.999)
        x = Reference(ws, min_col=1, min_row=first, max_row=last)
        y = Reference(ws, min_col=9, min_row=first, max_row=last)
        s = Series(y, x, title="kW")
        s.marker.symbol = "none"
        s.smooth = False
        sc.series.append(s)
        style_chart(sc, f"{beschr} - 15-Minuten-Leistung (kW)", [GRUEN], linie=True, breite=9525)
        wv = wb.create_sheet(f"Verlauf{suffix}")
        wv.add_chart(sc, "A1")

        # Ergänzung: Tagesenergie als Säulen - bei langen Zeiträumen deutlich lesbarer
        z1, z2 = summen_zeilen[obis]
        bt = BarChart()
        bt.type, bt.gapWidth = "col", 40
        bt.height, bt.width = 12.5, 34
        bt.legend = None
        bt.add_data(Reference(st, min_col=3, min_row=z1, max_row=z2), titles_from_data=False)
        bt.set_categories(Reference(st, min_col=6, min_row=z1, max_row=z2))
        bt.x_axis.tickLblSkip = 1
        bt.x_axis.majorTickMark = "none"
        style_chart(bt, f"{beschr} - Energie je Tag (kWh)", [GRUEN], linie=False)
        bt.x_axis.txPr = _txt_schraeg(900)
        wv.add_chart(bt, "A28")

        # Tagesvergleich: Diagramm ganz links, Werte rechts (ab Spalte V).
        # Eine Zeile je Tag -> der AutoFilter auf "Datum" blendet Tage samt Linie/Legende ein/aus
        # (Excel zeichnet ausgeblendete Zeilen nicht mit).
        wc = wb.create_sheet(f"Tagesvergleich{suffix}")
        T0, KOPF, ERSTE = 22, 4, 5                       # Spalte V, Kopfzeile 4, Daten ab Zeile 5
        zeiten = [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 15, 30, 45)]
        lookup: dict[date, dict[str, float]] = defaultdict(dict)
        for i in zeilen:
            lookup[ws.cell(i, 3).value].setdefault(ws.cell(i, 4).value, ws.cell(i, 9).value)  # DST: erster gewinnt
        wc.cell(1, T0, "Tage ein-/ausblenden: Pfeil bei „Datum“ (V4) anklicken, Haken setzen oder entfernen.").font = Font(bold=True)
        wc.cell(2, T0, f"Beim Öffnen sind nur die letzten {SICHTBAR_TAGE} Tage sichtbar. "
                       "Im Filter „Alles auswählen“ zeigt alle Tage; der Baum (Jahr > Monat > Tag) blendet ganze Monate auf einmal aus.").font = Font(italic=True, color="7F7F7F")
        wc.cell(3, T0, "X-Werte")
        wc.cell(KOPF, T0, "Datum")
        for j, z in enumerate(zeiten):
            wc.cell(3, T0 + 1 + j, j * 0.25 / 24).number_format = "hh:mm"  # X-Werte (Tagesbruchteil) fürs Diagramm
            wc.cell(KOPF, T0 + 1 + j, z)
        for c in range(T0, T0 + 1 + len(zeiten)):
            kopf = wc.cell(KOPF, c)
            kopf.fill, kopf.font = HEAD_FILL, HEAD_FONT
            kopf.alignment = Alignment(horizontal="center")
            wc.cell(3, c).font = Font(color="A6A6A6", size=8)
        sichtbar = set(tage[-SICHTBAR_TAGE:])
        for k, t in enumerate(tage):
            r = ERSTE + k
            wc.cell(r, T0, t).number_format = DATE_FMT.replace("%d", "DD").replace("%m", "MM").replace("%Y", "YYYY")
            for j, z in enumerate(zeiten):
                v = lookup[t].get(z)
                if v is not None:
                    wc.cell(r, T0 + 1 + j, v).number_format = "0.00"
            if t not in sichtbar:
                wc.row_dimensions[r].hidden = True
        ende = ERSTE + len(tage) - 1
        wc.auto_filter.ref = f"{get_column_letter(T0)}{KOPF}:{get_column_letter(T0)}{ende}"
        wc.auto_filter.filterColumn.append(FilterColumn(colId=0, filters=Filters(dateGroupItem=[
            DateGroupItem(year=t.year, month=t.month, day=t.day, dateTimeGrouping="day")
            for t in sorted(sichtbar)])))
        wc.column_dimensions[get_column_letter(T0)].width = 14
        for j in range(len(zeiten)):
            wc.column_dimensions[get_column_letter(T0 + 1 + j)].width = 6.5

        lc = ScatterChart()
        lc.height, lc.width = 15, 34
        xw = Reference(wc, min_col=T0 + 1, max_col=T0 + len(zeiten), min_row=3)
        for k in range(len(tage)):
            r = ERSTE + k
            se = Series(Reference(wc, min_col=T0 + 1, max_col=T0 + len(zeiten), min_row=r), xw)
            se.tx = SeriesLabel(strRef=StrRef(f"'{wc.title}'!${get_column_letter(T0)}${r}"))
            se.marker.symbol, se.smooth = "none", False
            lc.series.append(se)
        lc.x_axis.scaling.min, lc.x_axis.scaling.max = 0, 1
        lc.x_axis.majorUnit = 2 / 24                  # Beschriftung alle 2 h als Uhrzeit
        lc.x_axis.number_format = "[hh]:mm"
        lc.visible_cells_only = True          # ausgeblendete Zeilen werden nicht gezeichnet
        style_chart(lc, f"{beschr} - Tage im Vergleich (kW)",
                    [PALETTE[k % len(PALETTE)] for k in range(len(tage))], linie=True,
                    breite=15875, legende="r", plot_breite=0.80)
        wc.add_chart(lc, "A1")
        wc.sheet_view.zoomScale = 90

    add_vergleichsblaetter(wb, rows, csv_ist_kwh)

    # Reihenfolge: Daten, Tagessummen, Monats-/Wochenvergleich, übrige Diagramme, Listen zuletzt
    order = {"Daten": 0, "Tagessummen": 1, "Monatsvergleich": 2, "Wochenvergleich": 3}
    wb._sheets.sort(key=lambda w: next((v for k, v in order.items() if w.title.startswith(k)), 4))
    wb.save(out)


# --------------------------------------------------------------------------- main

def parse_day(s: str) -> date:
    return datetime.strptime(s, DATE_FMT).date()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0],
                                 formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--von", type=parse_day, help="erster Tag, TT.MM.JJJJ")
    ap.add_argument("--bis", type=parse_day, help="letzter Tag (Standard: = --von)")
    ap.add_argument("--update", action="store_true",
                    help=f"fehlende Tage ab {FIRST_DAY.strftime(DATE_FMT)} (bzw. --von) bis gestern (bzw. --bis)\n"
                         "nachladen; bereits vorhandene Tage werden übersprungen")
    ap.add_argument("--anlage", default=DEFAULT_ANLAGE)
    ap.add_argument("--zaehlpunkt", default=DEFAULT_ZAEHLPUNKT)
    ap.add_argument("--csv-dir", type=Path, default=Path("csv"),
                    help="Ordner für CSVs (Download-Ziel bzw. Quelle mit --nur-excel)")
    ap.add_argument("--nur-excel", action="store_true",
                    help="kein Download, nur vorhandene CSVs in --csv-dir verarbeiten")
    ap.add_argument("--out", type=Path, help="Excel-Datei (Standard: viertelstunden_<von>_<bis>.xlsx)")
    ap.add_argument("--headed", action="store_true", help="Browser sichtbar (Debug)")
    ap.add_argument("--csv-ist-kwh", action="store_true",
                    help="CSV-Werte sind echte kWh je 15 min (kW = Wert*4). Standard: Wert = kW,\n"
                         "wie im Portal-Diagramm")
    ap.add_argument("--version", action="version", version=__version__)
    a = ap.parse_args()

    if a.update:
        von = a.von or FIRST_DAY
        bis = a.bis or (date.today() - timedelta(days=1))
        tage = fehlende_tage(a.csv_dir, von, bis) if a.csv_dir.is_dir() else list(daterange(von, bis))
        if tage:
            print(f"{len(tage)} fehlende Tage zwischen {von.strftime(DATE_FMT)} und "
                  f"{bis.strftime(DATE_FMT)} - lade nach ...")
            download_csvs(tage, a.anlage, a.zaehlpunkt, a.csv_dir, a.headed, STOP_AFTER_EMPTY)
        else:
            print("Alles aktuell - keine fehlenden Tage.")
        files = sorted(a.csv_dir.glob("*.csv"))
    elif a.nur_excel or a.von is None:
        files = sorted(a.csv_dir.glob("*.csv"))
        if not files:
            ap.error(f"keine CSVs in {a.csv_dir} - --von/--bis angeben für Download")
    else:
        bis = a.bis or a.von
        if bis < a.von:
            ap.error("--bis liegt vor --von")
        print(f"Lade {a.von.strftime(DATE_FMT)} bis {bis.strftime(DATE_FMT)} ...")
        files = download_csvs(list(daterange(a.von, bis)), a.anlage, a.zaehlpunkt,
                              a.csv_dir, a.headed)

    rows = load_all(files)
    if not rows:
        print("Keine Daten gefunden.", file=sys.stderr)
        return 1
    first, last = rows[0]["start"], rows[-1]["start"]
    out = a.out or Path(f"viertelstunden_{first:%Y%m%d}_{last:%Y%m%d}.xlsx")
    build_excel(rows, out, a.csv_ist_kwh)
    print(f"{len(rows)} Werte -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
